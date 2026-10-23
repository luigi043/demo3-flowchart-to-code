"""Build the report from captured outputs; run with bundled Python."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image

base = Path(__file__).resolve().parent
evidence = base / 'test-evidence'
results = json.loads((evidence / 'results.json').read_text())
results += json.loads((evidence / 'browser-results.json').read_text())
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.add_paragraph('Flowchart to Code Test Evidence', 'Title')
doc.add_paragraph('Introduction to Software Engineering Demo 3', 'Subtitle')
doc.add_paragraph('Verification date 23 October 2026')
doc.add_paragraph('The converter turns a visual even or odd algorithm into matching pseudocode, Python and browser JavaScript. The automated checks below verify the conversion pipeline, generated program behavior, API error handling and production builds.')
doc.add_heading('Verification outcome', level=1)
doc.add_paragraph('All four command checks passed with exit code 0: backend compilation, 16 backend tests, 2 frontend tests and the frontend production build. All five browser checks passed, with no browser page errors.')
doc.add_heading('Evidence coverage', level=1)
doc.add_paragraph('Backend verbose output names every pipeline and behavior test. The generated JavaScript runs with a controlled browser prompt; the generated Python runs in an actual Python interpreter. Both cover negative even input, zero and odd input. Additional checks cover shared branch continuation, invalid graphs and preservation of API validation messages.')
doc.add_paragraph('The images include browser screenshots of captured command transcripts and direct screenshots of the running editor. Original PNG files, HTML transcripts and raw logs are preserved in docs/test-evidence. Standard output precedes standard error in each saved transcript; this can change display ordering without changing the recorded results.')
doc.add_heading('Manual checks and limitations', level=1)
doc.add_paragraph('The browser checks verified all three output tabs, clipboard content, JavaScript file download, label editing, deletion and the resulting validation message, with no page errors. Node dragging and connection editing were not exercised. The converter supports numeric input and structured acyclic flowcharts; loops and language specific expressions are outside the lesson scope.')
doc.add_paragraph('An earlier sandbox run could not write the Jest cache. A subsequent permitted run passed. A frontend test cleanup callback initially failed type checking; the callback was corrected and the final captured build below verifies the fix.')
doc.add_paragraph('The first browser copy assertion compared Windows clipboard CRLF text with LF text. The test harness now normalizes line endings; the subsequent copy and download checks passed. The initial browser transcript is retained as browser-smoke-initial.log.')
for i, r in enumerate(results, 1):
    doc.add_page_break()
    doc.add_heading(r['title'], level=1)
    doc.add_paragraph(r['command'], style='Normal')
    filename = evidence / f"{r['id']}.png"
    w, h = Image.open(filename).size
    width = min(6.8, 8.05 * w / h)
    doc.add_picture(str(filename), width=Inches(width))
    kind = 'Browser check' if r['id'].startswith('ui-') else 'Command result'
    doc.add_paragraph(f"Figure {i}. {kind}: {'PASS' if r['exitCode'] == 0 else 'FAIL'}. Source: {filename.name}.", style='Caption')
doc.save(base / 'Test Evidence.docx')
print(base / 'Test Evidence.docx')
